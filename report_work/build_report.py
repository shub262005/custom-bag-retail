from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
import hashlib, io, json, re, math
from lxml import etree as E
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches

ROOT=Path(__file__).resolve().parent
SOURCE=Path('D:/Submission/Project_Design_Report_Format_2026_27.docx')
OUT=ROOT.parent/'Roopam_Bag_Mall_Project_Design_Report_2026_27.docx'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS={'w':W,'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
def q(t): return '{'+W+'}'+t

# Diagrams are original code-drawn representations of the repository design.
FONT='C:/Windows/Fonts/times.ttf'
BOLD='C:/Windows/Fonts/timesbd.ttf'
def canvas(h):
 im=Image.new('RGB',(1560,h),'white'); return im,ImageDraw.Draw(im)
def label(d,xy,text,size=29,bold=False,anchor='mm'):
 d.multiline_text(xy,text,font=ImageFont.truetype(BOLD if bold else FONT,size),fill='black',anchor=anchor,align='center',spacing=7)
def box(d,xy,text,size=29):
 d.rectangle(xy,outline='black',width=3); label(d,((xy[0]+xy[2])/2,(xy[1]+xy[3])/2),text,size)
def arrow(d,a,b,text='',pos=None,size=26):
 d.line([a,b],fill='black',width=3)
 angle=math.atan2(b[1]-a[1],b[0]-a[0]); n=14
 pts=[b,(b[0]-n*math.cos(angle-.45),b[1]-n*math.sin(angle-.45)),(b[0]-n*math.cos(angle+.45),b[1]-n*math.sin(angle+.45))]
 d.polygon(pts,fill='black')
 if text:
  at=pos or ((a[0]+b[0])/2,(a[1]+b[1])/2-22)
  bounds=d.multiline_textbbox(at,text,font=ImageFont.truetype(FONT,size),anchor='mm',align='center',spacing=7)
  d.rectangle((bounds[0]-4,bounds[1]-3,bounds[2]+4,bounds[3]+3),fill='white')
  label(d,at,text,size)
def save(im,n):
 p=ROOT/(n+'.png'); im.save(p); return p

im,d=canvas(1440)
box(d,(40,30,730,130),'Customer / public visitor',33)
box(d,(830,30,1520,130),'Cashier / inventory manager / admin',33)
arrow(d,(385,130),(385,230),'Browse / design / track',(385,180))
arrow(d,(1175,130),(1175,230),'Sell / receive / review',(1175,180))
box(d,(40,230,730,405),'React customer interface\nStorefront | designer | request history')
box(d,(830,230,1520,405),'React staff interface\nPOS | inventory | purchases | reports')
box(d,(40,475,730,595),'Three.js / React Three Fiber\nLocal 3D rendering from design configuration')
arrow(d,(385,405),(385,475),'Render',(465,440))
box(d,(40,700,1520,835),'Spring Boot REST API /api/v1\nJWT filter + role rules → controllers + Jakarta validation',32)
d.line([(730,380),(780,380),(780,650)],fill='black',width=3)
arrow(d,(780,650),(780,700),'JSON / multipart + bearer JWT',(1015,665),28)
arrow(d,(1175,405),(1175,700),'JSON + JWT',(1315,540),28)
box(d,(40,930,735,1085),'Retail services\nCatalog | stock ledger | sale / purchase\nPayments | audit | reports',29)
box(d,(825,930,1520,1085),'Custom Bag services\nModel rules | server pricing | ownership\nReview state machine | logo handling',29)
arrow(d,(385,835),(385,930),'Validated commands',(385,885))
arrow(d,(1175,835),(1175,930),'Validated commands',(1175,885))
box(d,(40,1190,1040,1340),'PostgreSQL through JPA repositories\nUsers | catalog | sales | purchases | ledger | requests\nFlyway V1–V11 schema migrations',29)
box(d,(1120,1190,1520,1340),'Private logo files\nOwner/admin access',29)
arrow(d,(385,1085),(385,1190),'Transactional read / write',(385,1138))
arrow(d,(1010,1085),(1010,1190),'Persist',(940,1145))
arrow(d,(1320,1085),(1320,1190),'Store / load',(1320,1138))
label(d,(780,1395),'Responses return through the API to previews, status views, stock tables and reports. No external service is required.',25)
architecture=save(im,'architecture')

im,d=canvas(1050)
xs=[155,555,980,1400]
for x,t in zip(xs,['Customer UI','Security / controller','Request service','Database / files']):
 box(d,(x-140,30,x+140,125),t,26); d.line((x,125,x,990),fill='#888888',width=2)
steps=[(0,1,'1  POST multipart design + optional logo + JWT'),(1,2,'2  Authenticate CUSTOMER and validate DTO'),(2,3,'3  Load authenticated owner; validate / store logo'),(2,2,'4  Enforce bag rules; recompute estimate'),(2,3,'5  Save SUBMITTED request; flush transaction'),(3,2,'6  Saved entity / persistence error'),(2,1,'7  Response DTO or safe error'),(1,0,'8  Show request number, estimate and status')]
for j,(a,b,t) in enumerate(steps):
 y=205+j*97
 if a==b:
  d.line([(xs[a],y),(xs[a]+110,y),(xs[a]+110,y+35),(xs[a],y+35)],fill='black',width=3)
  label(d,(790,y-26),t,25)
 else: arrow(d,(xs[a],y),(xs[b],y),t,((xs[a]+xs[b])/2,y-25),25)
label(d,(780,1008),'If database persistence fails after file storage, the service deletes the newly stored logo.',26)
sequence=save(im,'sequence')

im,d=canvas(1300)
box(d,(40,35,490,190),'categories\nPK category_id | name | status',28)
box(d,(1070,35,1520,190),'brands\nPK brand_id | name | status',28)
box(d,(545,275,1015,480),'products\nPK product_id | SKU | prices\nFK category_id | FK brand_id\nstock_quantity | minimum_stock',28)
arrow(d,(265,190),(545,300),'1 : many',(330,245))
arrow(d,(1295,190),(1015,300),'1 : many',(1230,245))
box(d,(40,580,490,760),'sale_items\nPK sale_item_id\nFK sale_id | FK product_id\nquantity | unit price | line total',27)
box(d,(1070,580,1520,760),'purchase_items\nPK purchase_item_id\nFK purchase_id | FK product_id\nquantity | unit price | line total',27)
arrow(d,(610,480),(375,580),'1 : many',(420,510))
arrow(d,(950,480),(1185,580),'1 : many',(1160,510))
box(d,(545,580,1015,760),'inventory_transactions\nPK transaction_id\nFK product_id | type | quantity\nreference | date',27)
arrow(d,(780,480),(780,580),'1 : many',(860,530))
box(d,(40,865,490,1040),'sales\nPK sale_id | sale number | status\n1 : 1 payment | 1 : many audit rows',27)
box(d,(1070,865,1520,1040),'purchases\nPK purchase_id | FK supplier_id\n1 : many supplier payments',27)
arrow(d,(265,865),(265,760),'1 : many',(350,813))
arrow(d,(1295,865),(1295,760),'1 : many',(1380,813))
box(d,(545,865,1015,1040),'users\nPK user_id | unique email\npassword_hash | role | status',28)
box(d,(340,1140,1220,1270),'custom_bag_requests\nPK custom_bag_request_id | FK customer_user_id | unique request_number\nbag options | estimated_price | status | logo_reference | admin_note',27)
arrow(d,(780,1040),(780,1140),'1 : many',(865,1090))
schema=save(im,'schema')

im,d=canvas(1400)
label(d,(780,35),'DESIGN WIREFRAMES   Schematic layouts derived from implemented routes',30,True)
def frame(y,title):
 box(d,(40,y,1520,y+365),'',28); d.rectangle((40,y,1520,y+55),outline='black',width=2); label(d,(780,y+28),title,29,True)
frame(90,'Customer designer   /custom-bag')
box(d,(65,170,500,400),'Model / size / material\nBody / pocket / strap colours\nCompartments / supported options\nLogo / text / customer notes',26)
box(d,(535,170,1060,400),'Interactive 3D preview\nRotate / zoom\nBackpack | Laptop | Duffel',28)
box(d,(1095,170,1495,400),'Estimated price\nReview configuration\nSubmit request',28)
frame(525,'Staff point of sale   /pos')
box(d,(65,605,720,835),'Product search / selection\nProduct | price | available stock\nAdd product to cart',29)
box(d,(755,605,1495,835),'Cart items | quantity | unit price\nDiscount | payment method | total\nValidate and complete sale',29)
frame(960,'Admin request review   /custom-bag-requests')
box(d,(65,1040,500,1270),'Filter status / number / customer\nRequest list\nSelect request',27)
box(d,(535,1040,1060,1270),'Saved 3D reconstruction\nOptions / logo / customer notes\nEstimate and ownership',27)
box(d,(1095,1040,1495,1270),'Store note\nAllowed next status\nSave review',28)
label(d,(780,1360),'Wireframes describe interface structure; they are not captured screenshots or a separate implemented UI.',26)
wireframe=save(im,'wireframe')

# Values below map to original body child indices, recorded in artifact.md.
C={
5:'“Roopam Bag Mall Retail and Custom Bag Management System”',
10:'Under the Guidance of: [Insert Guide Name]',
14:'This is to certify that the Project Design Report entitled “Roopam Bag Mall Retail and Custom Bag Management System” has been prepared by [Insert Student Names] under the guidance of [Insert Guide Name] as part of Major Project – Phase I, Department of Information Technology, Walchand Institute of Technology, Solapur, during academic year 2026–2027.',
18:[
'Roopam Bag Mall Retail and Custom Bag Management System is a web application that connects the daily operations of a bag retailer with a structured customer customisation process. Separate stock records, sales records and informal design discussions can make it difficult to trace stock changes or preserve the exact configuration requested by a customer. The project addresses these workflow needs through a shared catalog, inventory ledger, purchases, counter sales, reports and customer design requests.',
'The architecture uses a React and TypeScript interface, a Spring Boot REST backend and a PostgreSQL database. Three.js and React Three Fiber render configurable Backpack, Laptop Bag and Duffel Bag models. The server validates model-specific options, derives request ownership from the authenticated identity, recalculates the estimated price and stores the design for later reconstruction. Administrators review requests through controlled status transitions. Retail transactions update stock and retain transaction and audit information.',
'A functioning local prototype is available. The integration record dated 5 October 2026 reports 317 passing backend tests with PostgreSQL regression suites enabled and a successful frontend build. These are recorded prototype results rather than evidence of production deployment. The remaining work includes formal user evaluation, production configuration, backup validation and presentation preparation. Custom bag submissions are enquiries for store review; they do not reserve inventory, collect payments or create retail orders.'
],
20:'We express our gratitude to [Insert Guide Name] for guidance during the design and development of this project. We thank the Department of Information Technology, the project coordinator, the Head of Department and Walchand Institute of Technology, Solapur, for academic support and facilities. We acknowledge the authors of the research cited in this report and the maintainers of the open-source technologies used in the prototype. We also thank our families and peers for their encouragement.',
28:[
'Bag retail involves maintaining product variants, monitoring stock, receiving purchases from suppliers, recording counter sales and answering customer enquiries. A customer may also want a bag with a particular material, colour, compartment layout, logo or text. These operations share business information but require different permissions and different levels of confirmation.',
'The project provides one browser application for customers and store staff. Customers browse the public storefront and submit saved custom bag designs. Cashiers operate the point of sale, inventory managers maintain stock and procurement, and administrators review custom requests. A local Java backend and relational database provide the common processing and storage layer.'
],
30:'The motivation is to make retail records and customer design communication consistent and traceable. A spreadsheet can represent quantities, but coordinating sales, purchase reversals, permissions and saved design options requires additional controls. A static product image also cannot show the effect of a customer’s chosen colours and supported options. This project therefore combines transaction-based stock handling with interactive 3D configuration and a reviewable request record. Reduction in staff effort and improvement in customer understanding are evaluation goals; no field-study improvement percentage is claimed.',
32:'Roopam Bag Mall Retail and Custom Bag Management System combines a public storefront, customer registration, a 3D bag designer, customer request tracking, administrator review, catalog management, inventory history, supplier management, purchases, POS, sales history and sales reports. The customer interface and staff interface use the same secured API while presenting navigation appropriate to each role. Standard custom bag requests preserve the chosen configuration and store response separately from retail sales.',
34:[
'Aim: To design and implement an integrated web prototype for traceable bag retail operations and customer custom bag requests.',
'O-01: Maintain products, categories, brands and suppliers with validated identifiers and active/inactive states.',
'O-02: Record purchases, sales and permitted reversals so that resulting stock quantities agree with their inventory movements in regression tests.',
'O-03: Enforce the CUSTOMER, CASHIER, INVENTORY_MANAGER and ADMIN permission matrix at the backend, including request ownership checks.',
'O-04: Allow customers to configure all three supported bag models, review an estimate, submit a standard request and reopen its saved 3D representation.',
'O-05: Permit only valid administrator review transitions and expose the saved status and store note to the owning customer.',
'O-06: Validate the prototype through automated backend tests, frontend build checks, workflow review and planned task-based usability testing.'
],
36:[
'Included: the application architecture, relational schema, role matrix, transaction workflows, rule-based bag configuration, price estimation, private logo handling, screen design and prototype verification. The design supports one store and its local demonstration environment.',
'Outside the present scope: an online checkout, payment gateway, delivery tracking, custom request stock reservation, manufacturing automation, arbitrary free-form bag geometry, SPECIAL_DESIGN processing, AI recommendations and demand forecasting. Logo storage is local, internet connectivity is required for API operations, and the 3D interface depends on a WebGL-capable browser. Production load capacity and measured usability remain to be evaluated.'
],
39:[
'The review covers ten papers published between 2021 and 2025, selected for product configuration, visual customisation, retail interaction and inventory design. Domain differences are stated explicitly. The papers motivate design choices; their reported benefits are not treated as measured results for this project.',
'[1] Yi, Huang and Yu (2022) compared navigation and pricing presentations in a mobile sneaker configurator with 180 students. Their experiment found that navigation and pricing combinations affect usability and perceived complexity. The narrow age group and footwear context limit transfer to bag retail. The implication here is to keep options accessible and make the estimate visible, then evaluate the actual bag interface.',
'[2] Dudek, Patalas-Maliszewska and Kowalczewska (2023) integrated customer preferences, expert rules and ERP information in a mattress manufacturer’s order configuration process. The case demonstrates rule-guided translation of preferences into feasible choices. Its manufacturing-specific knowledge base differs from this store prototype. The relevant design decision is to validate bag-type constraints on the server and retain human review.',
'[3] Li et al. (2023) proposed a social manufacturing framework combining multimodal information, 3D generation, blockchain and cloud management. The work provides an architectural account of connecting customisation to production services. It does not establish a small-store bag configuration benchmark. For this project, explicit configuration parameters provide a more bounded implementation than generated geometry.',
'[4] Manavis et al. (2024) developed a furniture customisation approach using parametric CAD and automated downstream artifacts, including assembly information. The case shows how user parameters can drive a consistent digital model. Rigid furniture geometry and fabrication outputs differ from sewn bags. The project adopts parameter-driven reconstruction while leaving manufacturing files outside scope.',
'[5] Grosso et al. (2021) examined 105 active online sales configurators and 522 Italian agrifood SME websites. The study identifies opportunities for configurators to improve the customer experience of smaller firms. It is exploratory and domain-specific rather than an operational bag-store trial. Its relevance is to connect a limited configuration space to an achievable SME workflow.',
'[6] Lu and Hsiao (2022) modelled virtual shopping adoption through diffusion of innovations, spatial presence, flow and satisfaction. Their empirical study links satisfaction and flow to purchase intention. Behavioural intention in a VR shopping context does not prove sales improvement in a browser bag designer. The project therefore treats visual engagement as a usability question and avoids a claimed conversion uplift.',
'[7] Li et al. (2024) analysed a two-echelon hybrid manufacturing supply chain using a Stackelberg game. Their results show that suitable customisation degree depends on cost structures and channel preferences. The analytical setting assumes manufacturing and market conditions that are not measured here. The practical implication is to label the calculated bag price as an estimate and retain store confirmation.',
'[8] Carpitella and Izquierdo (2025) reviewed 52 papers on sustainable inventory management and Industry 4.0. The review connects digital records and analytics with resource-management objectives. It synthesises heterogeneous studies rather than proving a specific retailer’s savings. The present design first establishes a stock ledger and reliable data; environmental and forecasting claims require later measurements.',
'[9] Naldi et al. (2025) mapped mass customisation research across conceptual metrics, enabling technologies and supply-chain integration. The review highlights modularity, customer participation and integration challenges. Its broad synthesis does not substitute for local feasibility or user testing. The project consequently bounds supported variants and separates a customer request from a committed manufacturing order.',
'[10] Desai (2025) described a Progressive Web Application for inventory work, including scanning, cross-device use and offline capabilities. The paper is an arXiv preprint and its development discussion is not independent proof of scalability. It motivates responsive access and identifies offline inventory as future work; the present React application does not claim PWA or offline synchronisation support.'
],
41:'Table 2.1 compares the reviewed approaches. Limitations describe applicability to the present project and do not imply that each paper intended to solve the same retail problem.',
44:[
'A manual register or spreadsheet workflow typically receives goods, records quantities, bills a sale and reconciles stock separately. Customer designs may be described through messages or sketches, which leaves the final specification dependent on interpretation. This is a baseline workflow model for analysis, not a claim that a formal on-site process audit has been completed.',
'An inventory or POS application focuses on catalog, stock, purchasing and billing. A visual configurator focuses on selecting and displaying product variants. The research reviewed above explores these areas in different contexts. The project brings a bounded visual enquiry workflow into the same application as store operations, while keeping its data lifecycle distinct from sales. No feature-by-feature commercial product benchmark is claimed.'
],
46:'The reviewed work suggests four design implications. First, navigation and price presentation require evaluation [1], so the designer uses grouped options and a review step. Second, configuration feasibility needs explicit rules [2], [4], so unsupported model options are rejected by the server. Third, operational integration is a continuing challenge [8], [9], so retail stock movements and saved custom requests share authentication and storage infrastructure but remain distinct transactions. Fourth, production cost and fulfilment uncertainty [7] justify administrator confirmation and an estimated price. This is a project-specific integration contribution, without a claim of a new algorithm.',
49:'A bag retailer needs consistent catalog, purchase, sales and stock records, while customers need a clear way to specify and revisit customised bag requirements. Disconnected records and informal descriptions make reconciliation and design review difficult. The project addresses these needs through a secured web system with transaction-linked stock updates, an interactive 3D designer and an owner-specific request review workflow.',
51:[
'Customer: registers, signs in, configures a standard bag, submits a request and reads only their own request details and store notes. Anonymous visitor: reads the public storefront and active brand/category advertising.',
'Cashier: reads product information and operates dashboard, POS and sales history/edit features. Inventory manager: additionally manages catalog, suppliers, inventory and purchases, accesses reports and cancels sales. Administrator: has staff capabilities and exclusive access to custom request review. Suppliers are business records rather than authenticated application users. The guide and evaluators review design evidence but do not require an application role.'
],
53:[
'FR-01: The system shall register a customer from validated name, email and password inputs, store a BCrypt hash and authenticate an active user to produce a JWT and role-aware session.',
'FR-02: The system shall return active brands and categories to public storefront visitors without exposing protected staff data.',
'FR-03: The system shall allow authorised managers/admins to create and update catalog and supplier records and reject invalid or duplicate identifiers.',
'FR-04: The system shall accept permitted stock movement inputs, validate quantities and product references, update stock and expose inventory history.',
'FR-05: The system shall save a purchase with supplier and line items, update stock, record supplier payments and reverse stock on valid cancellation.',
'FR-06: The system shall create and edit a counter sale, calculate item totals and discounts, record payment and stock changes, and retain audit information. Only manager/admin roles may cancel a sale.',
'FR-07: The system shall return date-based and grouped sales reports for authorised roles, keeping completed and cancelled semantics explicit.',
'FR-08: The system shall display Backpack, Laptop Bag and Duffel Bag configurations with supported options and a price estimate, then require review before submission.',
'FR-09: The system shall validate a multipart design request, derive its customer from the JWT, recompute pricing and return a saved request number. It shall accept only supported PNG/JPEG logos up to 2 MB.',
'FR-10: The system shall expose request history, reconstructed configuration and protected logos only to the owning customer or authorised administrator.',
'FR-11: The system shall accept only supported administrator review transitions, save the store note and return the updated status to the customer on retrieval.',
'FR-12: The system shall return structured validation, authentication, authorisation and not-found responses without exposing internal exception details.'
],
55:[
'NFR-01 Security: every protected endpoint shall enforce its role rule; all negative role and cross-owner scenarios in the regression matrix shall pass. Invalid/expired tokens shall be rejected and raw passwords shall never be returned.',
'NFR-02 Integrity and reliability: invalid retail operations shall roll back their database changes; verified purchase/sale reversal cases shall restore the expected stock. A failed request save shall trigger cleanup of its newly stored logo.',
'NFR-03 Performance target: on a documented local test machine with 10 concurrent users and 1,000 products, the 95th-percentile response time for catalog reads should be at most 2 seconds. This is a proposed acceptance target; no load-test result is available.',
'NFR-04 Usability target: at least 4 of 5 representative participants should complete a standard design submission and a staff sale task without assistance. Review the interface at 360, 768 and 1280 CSS-pixel widths, allowing horizontal scrolling inside staff data tables.',
'NFR-05 Privacy and operations: protected logos and customer notes shall remain inaccessible to another customer. Production readiness requires HTTPS, secret configuration and a documented database-plus-upload backup and restore trial; these deployment checks are planned.',
'NFR-06 Maintainability: schema changes shall use versioned Flyway migrations; the backend regression suite and frontend type/build checks shall pass before release. Larger catalog and request-list loads shall be measured before any scale claim.'
],
57:[
'UC-01 Configure a bag. Actor: CUSTOMER. Preconditions: active login and WebGL support. Main flow: choose model and options, inspect preview and estimate, review, submit, then reopen the saved request. Alternate flows: reject unsupported options or invalid logo; after a network failure, inspect request history before retrying to avoid duplicate submissions.',
'UC-02 Record a sale. Actor: CASHIER, INVENTORY_MANAGER or ADMIN. Preconditions: authorised session and valid product stock. Main flow: select products, enter quantities and discount, choose payment method, submit and inspect the saved sale. Invalid quantities or insufficient stock return an error.',
'UC-03 Review a request. Actor: ADMIN. Main flow: open inbox, filter requests, inspect saved configuration, choose an allowed next state and save a store note. A terminal-state reversal is rejected.',
'UC-04 Receive and cancel a purchase. Actor: INVENTORY_MANAGER or ADMIN. Main flow: choose supplier and items, save receipt and record payment; cancel only when reversal is valid. Dependencies: PostgreSQL availability, private upload storage and matching API deployment configuration. Formal approval of the design remains an academic review activity.'
],
60:'Figure 4.1 shows the project architecture. Browser interfaces exchange JSON or multipart requests with the Spring Boot API. Three-dimensional rendering runs in the browser. Business services persist relational data through JPA and store logo files through a private file service. No payment gateway, manufacturing service or external AI API participates in this flow.',
63:[
'Presentation layer: React Router separates public, customer and staff routes; protected and role routes control navigation. TanStack Query manages server data and Axios handles API communication. Customer and staff layouts present the functions appropriate to the signed-in user.',
'API and security layer: the JWT filter validates authentication, Spring Security applies endpoint permissions and controllers accept typed DTOs with Jakarta validation. CurrentUserService supplies trusted actor identity. Error handlers map failures to safe client responses.',
'Business layer: catalog services maintain master data; purchase and sale services coordinate amounts, payments, stock and reversals; report services aggregate sales; custom request services validate options, calculate estimates, enforce ownership and administer review states.',
'Persistence layer: PostgreSQL stores relational records and constraints, Flyway maintains schema versions, and Hibernate validates mapping compatibility. A separate local logo directory contains image data referenced by request records. Database and uploads together form the recoverable application state.'
],
65:[
'Customer data flow: (1) collect enumerated design options and optional branding; (2) render the current configuration locally; (3) submit a JSON request part and optional logo part; (4) validate role, fields and bag rules; (5) derive the customer identity and recompute the estimate; (6) store the request and logo reference; (7) return request details; (8) reconstruct the saved design on later retrieval.',
'Retail data flow: a purchase or sale DTO enters an authorised controller, service validation resolves product/supplier data, a transactional service updates the business document and stock ledger, and the response updates the browser view. Reports read persisted records with explicit date and status filters.',
'Control flow differs from payload flow: JWT and role checks decide whether processing may proceed; validation selects success or error paths; the review state machine selects permitted transitions; transaction commit/rollback determines persistence. A status change in a custom request never triggers a stock or payment operation.'
],
67:'A layered modular backend fits a single-store prototype and keeps transaction boundaries within one application. A relational database matches the many linked records and constraints in retail operations. Client-side parametric rendering gives immediate visual feedback without a rendering server. Server validation is retained because a browser can send modified requests. These choices respond to bounded configuration [2], [4], integration concerns [8], [9] and prototype resource limits. A microservice deployment would add operational work without a demonstrated need at this scale.',
70:'Table 5.1 identifies module boundaries and the main data exchanged. Controllers and security form shared dependencies. Validation failures stop processing; missing records return not-found responses; unexpected errors are logged on the server and represented safely to the client.',
73:'Figure 5.1 details the standard customer request interaction. Validation precedes persistence; the service obtains ownership from the signed-in identity rather than a customer identifier supplied by the browser. Logo cleanup compensates for a failed database save. Administrator review is a later authenticated request against the saved record.',
75:[
'Figure 5.2 shows the principal relational design. A category has many products, a brand can identify many products, and each sale/purchase contains multiple item rows linked to products. Supplier contacts and payments extend supplier and purchase records. Stock movements are represented by inventory transactions; sale audit history records relevant changes and actors.',
'A user owns many custom_bag_requests through customer_user_id. Each request stores enumerated model, size, material and colours, supported booleans, compartment count, optional logo/text metadata, estimated price, review status, notes and timestamps. request_number is unique. Logo bytes remain in protected file storage; the database keeps a generated reference.',
'Database constraints include primary/foreign keys, unique user email and request numbers, and checks on supported enum values and nonnegative estimates. Product SKU uniqueness and service validation protect catalog identity. Model-specific constraints are also checked in the service because permitted compartments and features vary by bag type.',
'The review workflow is SUBMITTED → REVIEWING → APPROVED → COMPLETED; SUBMITTED and REVIEWING may also become REJECTED. COMPLETED and REJECTED are terminal. CANCELLED exists in the schema but the current administrator transition map exposes no path into it. A stored request is not a sales order.'
],
77:[
'All application endpoints use /api/v1. JSON is the default representation; protected calls send Authorization: Bearer <token>. POST /auth/register and POST /auth/login are public; GET /auth/me returns authenticated user details. GET /public/storefront/brands and /categories provide public advertising records.',
'Catalog and retail interfaces: /products, /categories, /brands and /suppliers expose list/create and identifier-based detail/update/status operations. /inventory-transactions supports history and validated movement creation. /purchases supports receipt, edit, cancellation and nested payment operations. /sales supports creation, edit, detail, dashboard and cancellation. /sales/reports exposes daily, date-range, by-product, by-category, by-payment-method and cancelled views.',
'POST /custom-bag-requests accepts multipart/form-data with an application/json part named request and an optional image part named logo. The request specifies requestType, bagType, size, material, bodyColor, pocketColor, strapColor, functional options, compartmentCount, branding and customerNotes. The server response includes saved identifiers, configuration, estimatedPrice and status.',
'GET /custom-bag-requests/mine and /mine/{id} return owner-specific history and details; /mine/{id}/logo streams a protected logo. Administrator list/detail/logo paths use /custom-bag-requests/admin; PATCH /admin/{id} accepts status and adminNote. Cross-owner lookups return 404. Authentication failure returns 401, forbidden roles return 403, and validation failures return a structured client error. Hardware communication protocols are not applicable because this is a browser-based software system.'
],
79:[
'Access control is enforced twice: role-aware navigation helps users find permitted functions, and backend request rules independently enforce permissions. BCrypt protects stored passwords. JWT validation and current-user lookup establish identity; request ownership filters prevent a customer from reading another customer’s request or logo. The server, not the submitted payload, determines the review actor and request owner.',
'Jakarta constraints and service checks validate required fields, monetary and quantity rules, model-specific combinations and review transitions. Upload handling limits image type and size and uses generated file references. Unexpected exceptions are logged internally; client messages omit stack traces and internal exception text.',
'Transactional retail processing provides rollback on failures, while new-logo cleanup addresses file/database mismatch during request creation. Deployment hardening requires HTTPS, protected secrets, restricted filesystem permissions, backup/restore checks and a documented log-retention policy. An explicit concurrency/load review remains necessary before deployment; a passing functional regression does not establish production race safety.'
],
82:[
'The core method is deterministic, parameter-driven bag configuration combined with transactional retail logic. A finite configuration object controls browser geometry and materials. The backend checks that the object belongs to a supported solution space and calculates an authoritative estimate from fixed pricing rules. There is no training dataset, learned model or generative geometry pipeline.',
'Saved configuration values drive the same model components when a customer or administrator reopens a request. Review states constrain what the store may do next. Retail sale/purchase logic separately updates stock and monetary records so that a design enquiry cannot accidentally change store availability.'
],
84:[
'Algorithm 6.1 Standard custom request creation',
'Input: authenticated CUSTOMER, design DTO and optional logo. Output: saved request DTO or validation/error response.',
'1. Validate required enum fields, text lengths and supported logo constraints.',
'2. Check model rules: Backpack/Duffel allow 1–4 compartments; Laptop Bag allows 1–3. Reject side pockets for Laptop Bag and laptop padding for Duffel Bag. Reject SPECIAL_DESIGN.',
'3. Require logo position when a logo exists; require text colour and position when trimmed custom text is nonempty. Resolve the customer from the authenticated identity.',
'4. Store the valid logo under a generated reference. Calculate the estimate from the server rule table; do not accept a client-provided price as authoritative.',
'5. Generate the request number; populate configuration, owner, timestamps and SUBMITTED status. Save and flush within the service transaction.',
'6. If persistence raises an error, delete the newly stored logo and propagate the error. Otherwise return the saved DTO. On later reads, filter by owner or require ADMIN.',
'Price formulation: E = B + S + M + 100(c − 1) + 100f + 150s + 200l + 150w + 200g + 100t, in INR. Here c is compartment count and f, s, l, w, g, t indicate front pocket, side pockets, laptop padding, water resistance, logo and nonempty text respectively.',
'Base B: Backpack 1200, Laptop Bag 1400, Duffel Bag 1600. Size S: Small 0, Medium 150, Large 300. Material M: Polyester 0, Canvas 150, Leather 500. The result is represented to two decimal places. Example: a large Canvas Backpack with two compartments, side pockets, water resistance, logo and text has E = 1200 + 300 + 150 + 100 + 150 + 150 + 200 + 100 = INR 2350.00.',
'Algorithm 6.2 Retail stock processing: validate the authorised sale/purchase and line quantities; resolve active products and business rules; calculate line totals and permitted discounts; verify stock availability for outward movements; persist the document, associated payment/audit data and stock movement within a transaction. A permitted cancellation records the inverse stock effect and preserves the document’s cancelled state for reporting. Duplicate cancellation must be rejected.'
],
86:[
'Backend design: Spring controllers use typed DTOs, services express business rules and repositories persist JPA entities. Sale totals use decimal arithmetic; the current sale calculation subtracts the chosen discount from subtotal and does not add a separate tax amount. Product/category report amounts are gross item amounts, while completed-sale grand totals reflect sale-level discounts.',
'Frontend design: TypeScript models connect forms to the API. React Query refreshes server data, and session cleanup clears cached identity data after a 401. Three model components share rendering utilities but retain model-specific geometry. Saved colours, material, size, branding positions and options reconstruct the request.',
'Storage and recovery: Flyway V1–V11 define the prototype schema; Hibernate validates the mappings. Protected logo files and relational records must be migrated and backed up together. An annual sequence generates readable request numbers. No new dataset acquisition or AI training is required for this methodology.'
],
88:'The engineering contribution is the integration of a saved, reconstructable bag configuration and constrained administrator review into a retail management prototype with distinct role permissions and stock workflows. Compared with the different emphases of the reviewed studies, it combines a bounded customer interface, authoritative backend rules and an operational handover record in one application. The contribution is demonstrated through implemented modules and regression scenarios; improved user satisfaction, business revenue and manufacturing efficiency remain hypotheses for later evaluation.',
91:[
'Customer journey: /home → register or login → /custom-bag → review configuration → submit → /my-custom-bags → saved detail and store response. The designer allows supported model selection, preview rotation/zoom, option changes and optional logo/text. Submission requires a customer login.',
'Staff journey: login → role-appropriate /dashboard → catalog, inventory, purchases, POS, sales or reports. An administrator additionally opens /custom-bag-requests to review the same saved customer design. Cashiers do not receive report or cancellation controls; inventory managers do not receive the administrator request inbox. Forbidden access leads to an appropriate return route.'
],
93:'Figure 7.1 presents original schematic wireframes for the designer, POS and administrator review screens. They document the layout and actions of the implemented route families rather than claim to be screenshots. The designer groups configuration, 3D feedback and submission review; POS separates product selection from the cart; review presents stored design information alongside permitted status controls.',
95:[
'Functioning prototype: public storefront, role-aware authentication, three 3D bag models, standard request submission with optional logo, customer history/detail, administrator review, catalog, inventory, suppliers, purchases, POS, sales and reports are represented by implemented source modules.',
'Recorded integration review on 5 October 2026 exercised the customer and administrator saved preview workflow through APPROVED. Completion and terminal transition rejection were verified in transactional PostgreSQL tests; the persistent presentation example was not manually advanced to COMPLETED. The wireframes in this report are design representations. Production deployment, formal usability scoring and load benchmarks are planned.'
],
98:'Table 8.1 records technologies declared in the repository. Frontend versions are the dependency specifications, not a claim that every compatible patch has been evaluated. Selection emphasises typed interfaces, relational integrity and browser-based 3D interaction.',
101:'Table 8.2 separates implemented prototype modules from remaining release activities. Evidence is located in the repository and its integration record dated 5 October 2026. No fresh application test execution is implied by preparing this report.',
104:[
'The following four-week plan starts after the guide’s review; week numbers are relative because an approved submission calendar has not been supplied.',
'Week 1: review requirements and report with the guide, confirm genuine store contact details, complete team metadata and obtain representative user tasks. Deliverable: reviewed design and agreed acceptance checklist.',
'Week 2: conduct task-based usability and accessibility review, record task success/time, test responsive layouts and address confirmed defects. Deliverable: participant observations and prioritised fixes.',
'Week 3: configure a staging deployment with HTTPS and environment secrets; run load/concurrency checks and a database-plus-upload restore trial. Deliverable: measured performance and recovery evidence.',
'Week 4: rerun regression/build checks after fixes, collect approved demonstration screenshots, finalise report pagination and prepare the viva demonstration. Deliverable: reviewed release candidate, documentation and presentation. Deployment is conditional on successful acceptance checks.'
],
106:[
'Technical feasibility: the local prototype uses widely supported web and relational technologies, and the existing verification record supports functional feasibility. WebGL support, large frontend chunks and unmeasured concurrent writes remain practical constraints.',
'Economic feasibility: the chosen frameworks are open source and do not require a paid external API. Hosting, backup storage, maintenance and user-support costs still need a deployment budget; no cost-saving amount is claimed. Operational feasibility depends on staff entering accurate source records and reviewing custom requests before committing to manufacture.',
'Key risks and mitigations: stock races under concurrent requests require targeted concurrent transaction testing and locking review; loss of logo files requires coordinated database/file backups; oversized 3D bundles require load measurement and later optimisation; unverified store details require owner confirmation; unsupported option combinations require server validation; schedule risk is controlled by freezing the current scope before formal evaluation.'
],
109:[
'Unit tests cover pricing rules, validation, number generation, service operations and entity behaviour. Controller/security tests verify request mapping, safe errors and the permission matrix. PostgreSQL regressions run real migrations and transactional fixtures for authentication, sales, purchases and custom request persistence. Test fixtures must use an isolated development/test database, never a production database.',
'System testing follows customer design submission and staff retail journeys, including invalid inputs and forbidden roles. Planned usability evaluation measures task completion and assistance for at least five representative participants. Planned performance/recovery checks measure latency under stated load and restoration of both records and logos. Table 9.1 defines concrete acceptance scenarios.'
],
112:[
'Functional acceptance: all applicable automated cases pass without failures/errors; sale and purchase scenarios produce the expected ledger changes; owner/role restrictions reject every negative security case; saved requests reconstruct equivalent design values; invalid review transitions fail.',
'Recorded baseline: the 5 October 2026 integration record reports 317 tests passed with all opt-in PostgreSQL suites enabled. The default run reported 305 cases, of which 296 passed and 9 were skipped. The frontend lint run had zero errors and 15 warnings; the build passed with a large-chunk advisory. These dated results establish a baseline, not a current production certificate.',
'Release evaluation still requires the NFR-03 response-time target, NFR-04 task-success target and a demonstrated backup restore. No measured performance, SUS score or participant success rate is available in the current evidence.'
],
114:[
'Achieved prototype outputs: Spring Boot API, React/TypeScript interface, PostgreSQL migrations, three configurable 3D bag models, saved standard custom request workflow, staff retail modules and automated regression sources. Available design outputs include the architecture, interaction, schema and interface representations in this report.',
'Planned release outputs: reviewed user-evaluation results, load/concurrency evidence, staging deployment configuration, restore demonstration, approved screenshots, final presentation and completed personal/certificate metadata. A deployment URL and signed academic approval are not claimed.'
],
117:'The design connects customer customisation and store operations through a shared secured application. Parametric 3D models provide immediate visual feedback, server rules establish valid configurations and estimates, and saved owner-specific requests support review and later reconstruction. Transactional retail services maintain sale/purchase effects and stock history. The architecture and distinct workflows address traceability without treating an enquiry as a paid order.',
119:'The repository contains an implemented local prototype and a dated successful integration baseline. Academic approval of the final title and design remains with the guide and department. Immediate next steps are to insert team/guide details, review store information, complete user and performance evaluation, validate recovery and prepare the final demonstration. The existing schema, roles and module boundaries provide a basis for these activities; production readiness is not inferred from prototype completeness.',
121:'Future work may add offline staff operation with conflict resolution, barcode scanning, demand forecasting after sufficient quality data is available, multi-store stock transfers, automated customer notifications, accessibility improvements, configurable pricing administration and richer material previews. A separately designed order/quotation module could connect approved requests to payments and fulfilment. Such work requires explicit consent, security review, new business rules and testing rather than changing the meaning of existing request records.',
125:[
'Appendix A Detailed diagrams: Figure 4.1 defines deployment-layer responsibilities, Figure 5.1 traces standard request creation and Figure 5.2 records principal entity relationships. Sequence/number tables and supplier contacts are supporting schema objects and are omitted from the simplified ER view for readability.',
'Appendix B Supplementary requirements and schema: FR-01–FR-12 map to the acceptance tests in Table 9.1. The source schema is src/main/resources/db/migration, versions V1–V11. Relevant implementation evidence includes SecurityConfig.java, CustomBagPricingService.java, CustomBagRequestServiceImpl.java, SaleServiceImpl.java and PurchaseServiceImpl.java.',
'Appendix C Prototype evidence: frontend/src/features/custom-bag/three contains the designer and three model components. App.tsx defines routes and the features folders contain retail screens. README.md describes setup and roles. FINAL_INTEGRATION_REPORT.md, dated 5 October 2026, records the prototype review and its limitations. Figure 7.1 is a wireframe rather than a captured application image.',
'Appendix D Verification and work plan: target/final-all-tests.log, target/final-default-tests.log and target/surefire-reports contain the existing backend run evidence; target/final-build.log and target/final-lint.log record frontend checks. The four-week remaining plan appears in Section 8.3. Retain database and upload files together when moving the demonstration environment.'
]
}

refs=[
'[1] M. Yi, Z. Huang, and Y. Yu, “Creating a sustainable e-commerce environment: The impact of product configurator interaction design on consumer personalized customization experience,” Sustainability, vol. 14, no. 23, Art. no. 15903, 2022, doi: 10.3390/su142315903.',
'[2] A. Dudek, J. Patalas-Maliszewska, and K. Kowalczewska, “Automatic configuration of an order as an integral part of a cyber-physical system in a manufacturing operating according to mass-customisation strategy,” Applied Sciences, vol. 13, no. 4, Art. no. 2499, 2023, doi: 10.3390/app13042499.',
'[3] Y. Li et al., “An efficient product-customization framework based on multimodal data under the social manufacturing paradigm,” Machines, vol. 11, no. 2, Art. no. 170, 2023, doi: 10.3390/machines11020170.',
'[4] A. Manavis, P. Minaoglou, N. Efkolidis, and P. Kyratsis, “Digital customization for product design and manufacturing: A case study within the furniture industry,” Electronics, vol. 13, no. 13, Art. no. 2483, 2024, doi: 10.3390/electronics13132483.',
'[5] C. Grosso, F. Checchinato, V. Finotto, and C. Mauracher, “Configuration challenges for the ‘Made in Italy’ agrifood industry,” International Journal of Industrial Engineering and Management, vol. 12, no. 3, 2021. [Online]. Available: https://sp.ftn.uns.ac.rs/index.php/ijiem/article/view/63. [Accessed: Oct. 16, 2026].',
'[6] X. Lu and K.-L. Hsiao, “Effects of diffusion of innovations, spatial presence, and flow on virtual reality shopping,” Frontiers in Psychology, vol. 13, Art. no. 941248, 2022, doi: 10.3389/fpsyg.2022.941248.',
'[7] W. Li, H. Sun, M. Tong, N. Mustafee, and L. Koh, “Customizing customization in a 3D printing-enabled hybrid manufacturing supply chain,” International Journal of Production Economics, vol. 268, Art. no. 109103, 2024, doi: 10.1016/j.ijpe.2023.109103.',
'[8] S. Carpitella and J. Izquierdo, “Trends in sustainable inventory management practices in Industry 4.0,” Processes, vol. 13, no. 4, Art. no. 1131, 2025, doi: 10.3390/pr13041131.',
'[9] L. D. Naldi, F. G. Galizia, M. Bortolini, M. Gabellini, and E. Ferrari, “Unlocking the potential of mass customization through Industry 4.0: Mapping research streams and future directions,” Applied Sciences, vol. 15, no. 13, Art. no. 7160, 2025, doi: 10.3390/app15137160.',
'[10] A. Desai, “Enhancing inventory management with Progressive Web Applications (PWAs): A scalable solution for small and large enterprises,” arXiv preprint arXiv:2506.11011, 2025, doi: 10.48550/arXiv.2506.11011.'
]
C[123]=refs

tables={
9:[[str(i),'[Insert Name]','[Insert Roll No]','[Insert Email]','[Insert PRN]'] for i in range(1,5)],
42:[
['[1] / 2022','Mobile configurator experiment','Navigation/pricing affect usability','Student footwear setting','Visible options and estimate'],
['[2] / 2023','Rules with ERP knowledge','Feasible order configuration','Mattress-specific case','Server option validation'],
['[3] / 2023','Multimodal 3D framework','Connects design and services','Broad production framework','Bounded parameter model'],
['[4] / 2024','Parametric furniture CAD','Automated design outputs','Rigid furniture context','Saved model reconstruction'],
['[5] / 2021','SME configurator analysis','Customer experience opportunities','Agrifood exploratory study','Achievable retail scope'],
['[6] / 2022','VR adoption model','Flow/satisfaction relate to intent','VR intention, not bag sales','Evaluate preview usability'],
['[7] / 2024','Supply-chain game model','Costs shape customisation','Analytical assumptions','Estimate plus store review'],
['[8] / 2025','Inventory literature review','Digital data supports decisions','Heterogeneous evidence','Traceable stock ledger'],
['[9] / 2025','Mass customisation review','Modularity and integration matter','Broad synthesis','Finite options and workflow'],
['[10] / 2025','Inventory PWA preprint','Cross-device/offline design','Not independently validated','Responsive UI; offline later']
],
71:[
['Auth','Identity','Login / JWT','Hash and verify','User / token','Users / security'],
['Catalog','Master data','Product / SKU','Validate / save','Catalog records','Category / brand'],
['Stock','Trace stock','Movement DTO','Validate / adjust','Ledger / balance','Products'],
['Purchases','Receive goods','Supplier / items','Totals / receipt / reversal','Purchase / payments','Stock / suppliers'],
['Sales','Counter billing','Cart / discount','Totals / stock / audit','Sale / payment','Stock / actor'],
['Reports','Sales analysis','Dates / grouping','Filter / aggregate','Report DTOs','Sales / items'],
['Designer','Visual choices','Bag options','Render / review','3D / request DTO','Three.js / auth'],
['Requests','Save / review','DTO / logo / state','Validate / price / transition','Saved design / note','Users / files / DB']
],
99:[
['Programming Language','Java; TypeScript','21; ~6.0.2','Backend rules and typed browser code'],
['Framework / Library','Spring Boot; React; Vite','3.3.3; ^19.2.8; ^8.2.2','REST services and component interface'],
['Database / Storage','PostgreSQL; Flyway; local files','PostgreSQL 18.6 in recorded demo; V1–V11','Relational records, migrations and logos'],
['Tools / Platform','Maven; npm; Git; Tailwind CSS','Maven 3.9+ setup; Tailwind ^4.3.3','Builds, version control and styling'],
['Hardware / External API','WebGL browser; Three.js / Fiber / Drei','^0.186.1 / ^9.8.1 / ^10.7.9','3D preview; no external API required']
],
102:[
['Auth / role matrix','Completed prototype','SecurityConfig; auth and role tests','Review deployment secrets'],
['Catalog / suppliers / stock','Completed prototype','Services, migrations and screens','User acceptance review'],
['Purchases / sales / reports','Completed prototype','Service and PostgreSQL tests','Concurrency/load checks'],
['3D designer / saved requests','Completed prototype','Three model files; request tests','Usability evaluation'],
['Admin review states','Completed prototype','Transition map and tests','Confirm presentation sample'],
['Documentation','In progress','README and design report','Guide review and metadata'],
['Staging / recovery / load','Not started in supplied evidence','Planned Section 8.3 checks','Configure and measure'],
['Formal user evaluation','Not started in supplied evidence','Planned five-participant tasks','Recruit and record results']
],
110:[
['T-01','FR-01 Auth','Valid login; invalid/expired token','Login succeeds; bad tokens fail','Correct identity; 401 for invalid token'],
['T-02','FR-01 / 10 Roles','Cross-role API; another owner ID','403 role denial; owner 404','All negative role cases pass'],
['T-03','FR-03 Catalog','Duplicate SKU; invalid price','Reject invalid catalog request','No duplicate/invalid record'],
['T-04','FR-04 / 05 Stock','Purchase two units; cancel receipt','Stock increases then reverses','Expected ledger; original quantity'],
['T-05','FR-06 Sales','Create/edit sale; manager cancel','Adjust stock; inverse on cancel','Correct stock and audit actor'],
['T-06','FR-06 Security','Cashier requests sale cancellation','Forbidden; no mutation','403; stock unchanged'],
['T-07','FR-08 / 09 Pricing','Valid model; altered client estimate','Server calculates its price','Exact two-decimal rule total'],
['T-08','FR-09 Validation','Laptop side pockets; invalid logo','Reject unsupported input','No saved invalid request/file'],
['T-09','FR-10 Saved design','Submit then retrieve own request','Equivalent options and preview','All persisted values match'],
['T-10','FR-11 States','Allowed states; reopen terminal','Accept allowed; reject terminal','Correct state and saved note'],
['T-11','FR-07 Reports','Completed/cancelled sales by date','Correct filtered aggregates','Match fixtures and gross/net labels'],
['T-12','FR-12 Errors','Validation error; unexpected failure','Safe structured response','No stack trace/internal message'],
['T-13','NFR-03 Load','10 users; 1,000 products','Catalog p95 <= 2 seconds','Planned measured load trial'],
['T-14','NFR-04 Usability','5 users complete key tasks','At least 4 succeed unaided','Planned observed task success'],
['T-15','NFR-05 Restore','Restore database plus upload files','Saved requests/logos accessible','Planned record/file checks']
]
}

with ZipFile(SOURCE) as z: parts={n:z.read(n) for n in z.namelist()}
tree=E.fromstring(parts['word/document.xml']); body=tree.find('w:body',NS); orig=list(body)
body_sample=deepcopy(orig[28]); header_sample=deepcopy(orig[27])

def settext(p,text):
 # Preserve source paragraph and first run formatting without keeping guide text.
 run=p.find('w:r',NS); rp=deepcopy(run.find('w:rPr',NS)) if run is not None and run.find('w:rPr',NS) is not None else None
 for c in list(p):
  if c.tag!=q('pPr'): p.remove(c)
 r=E.SubElement(p,q('r'))
 if rp is not None:r.append(rp)
 for j,line in enumerate(text.split('\n')):
  if j:E.SubElement(r,q('br'))
  t=E.SubElement(r,q('t')); t.text=line; t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
 return p
def para(text,sample=None): return settext(deepcopy(sample if sample is not None else body_sample),text)
def replace(i,value):
 values=value if isinstance(value,list) else [value]
 settext(orig[i],values[0]); after=orig[i]
 for t in values[1:]:
  p=para(t,orig[i]); after.addnext(p); after=p
 return after
for i,t in C.items(): replace(i,t)

def fill_table(i,rows):
 tab=orig[i]; trs=tab.findall('w:tr',NS); sample=deepcopy(trs[1])
 for tr in trs[1:]:tab.remove(tr)
 # Keep the source table grid and cell properties exactly. Repeat header for multipage data.
 hp=trs[0].find('w:trPr',NS)
 if hp is None:hp=E.SubElement(trs[0],q('trPr'))
 E.SubElement(hp,q('tblHeader'))
 for row in rows:
  tr=deepcopy(sample)
  pr=tr.find('w:trPr',NS)
  if pr is None:pr=E.SubElement(tr,q('trPr'))
  E.SubElement(pr,q('cantSplit'))
  for cell,text in zip(tr.findall('w:tc',NS),row):
   ps=cell.findall('w:p',NS)
   settext(ps[0],text)
   for p in ps[1:]:cell.remove(p)
  tab.append(tr)
for i,rows in tables.items(): fill_table(i,rows)

# Caption tables in the source paragraph pattern; no new table styling.
for i,title in [(42,'Table 2.1 Comparative literature review'),(71,'Table 5.1 Module responsibilities and interfaces'),(99,'Table 8.1 Technology stack'),(102,'Table 8.2 Implementation status'),(110,'Table 9.1 Requirements and acceptance tests')]:
 p=para(title); orig[i].addprevious(p)
 pp=p.find('w:pPr',NS); E.SubElement(pp,q('keepNext'))

# Build drawing XML in a helper document; only additive image parts/relationships are imported.
helper=Document(); figures=[architecture,sequence,schema,wireframe]
for p in figures:helper.add_paragraph().add_run().add_picture(str(p),width=Inches(6.45))
b=io.BytesIO(); helper.save(b)
with ZipFile(io.BytesIO(b.getvalue())) as z:
 ht=E.fromstring(z.read('word/document.xml'))
 drawings=ht.findall('.//w:drawing',NS)
 hrels=E.fromstring(z.read('word/_rels/document.xml.rels'))
 newrels=E.fromstring(parts['word/_rels/document.xml.rels'])
 relns='http://schemas.openxmlformats.org/package/2006/relationships'
 rattr='{'+NS['r']+'}embed'
 for idx,drawing in enumerate(drawings):
  blip=drawing.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip'); rid=blip.get(rattr)
  target=next(x.get('Target') for x in hrels if x.get('Id')==rid)
  name='report_figure_'+str(idx+1)+'.png'; new_id='rIdReportFigure'+str(idx+1)
  E.SubElement(newrels,'{'+relns+'}Relationship',Id=new_id,Type=NS['r']+'/image',Target='media/'+name)
  parts['word/media/'+name]=z.read('word/'+target); blip.set(rattr,new_id)
  for el in drawing.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr'):
   el.set('id',str(100+idx));el.set('descr',['Overall retail and custom bag system architecture','Standard custom request sequence','Principal database entity relationships','Schematic designer POS and review screen wireframes'][idx])
parts['word/_rels/document.xml.rels']=E.tostring(newrels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=E.fromstring(parts['[Content_Types].xml']);ctns='http://schemas.openxmlformats.org/package/2006/content-types'
if not any(e.get('Extension')=='png' for e in ct):E.SubElement(ct,'{'+ctns+'}Default',Extension='png',ContentType='image/png')
parts['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)

figcaps=[]
def put_figure(after,drawing,title,replace_slot=False):
 p=after if replace_slot else para('')
 if not replace_slot:after.addnext(p)
 settext(p,''); r=p.find('w:r',NS);r.append(deepcopy(drawing))
 pp=p.find('w:pPr',NS);jc=pp.find('w:jc',NS)
 if jc is not None:jc.set(q('val'),'center')
 else:E.SubElement(pp,q('jc')).set(q('val'),'center')
 E.SubElement(pp,q('keepNext'))
 cap=para(title);p.addnext(cap);figcaps.append(cap)
 return cap
put_figure(orig[61],drawings[0],'Figure 4.1 Overall system architecture',True)
put_figure(orig[73],drawings[1],'Figure 5.1 Standard custom request creation sequence')
# Place schema after all prose belonging to 5.3.
after=orig[75]
while after.getnext() is not orig[76]:after=after.getnext()
put_figure(after,drawings[2],'Figure 5.2 Principal database entity relationships')
put_figure(orig[93],drawings[3],'Figure 7.1 Key screen wireframes')

# Stable bookmarks + PAGEREF fields populate the template's existing contents/list slots.
bookmark_id=100
def bookmark(p,name):
 global bookmark_id
 start=E.Element(q('bookmarkStart'));start.set(q('id'),str(bookmark_id));start.set(q('name'),name)
 end=E.Element(q('bookmarkEnd'));end.set(q('id'),str(bookmark_id));bookmark_id+=1
 p.insert(1 if p.find('w:pPr',NS) is not None else 0,start);p.append(end)
def pageref(p,name):
 r=E.SubElement(p,q('r'));E.SubElement(r,q('fldChar')).set(q('fldCharType'),'begin')
 r=E.SubElement(p,q('r'));E.SubElement(r,q('instrText')).text=' PAGEREF '+name+' \\h '
 r=E.SubElement(p,q('r'));E.SubElement(r,q('fldChar')).set(q('fldCharType'),'separate')
 r=E.SubElement(p,q('r'));E.SubElement(r,q('t')).text='0'
 r=E.SubElement(p,q('r'));E.SubElement(r,q('fldChar')).set(q('fldCharType'),'end')
heads=[26,37,47,58,68,80,89,96,107,115,122,124]
tocrows=orig[22].findall('w:tr',NS)[1:]
for i,(index,row) in enumerate(zip(heads,tocrows)):
 name='ReportChapter'+str(i+1);bookmark(orig[index],name)
 cell=row.findall('w:tc',NS)[-1];p=cell.find('w:p',NS);settext(p,'');pageref(p,name)
for listindex,prefix,items in [(23,'List of Figures',figcaps),(24,'List of Tables',[orig[i].getprevious() for i in [42,71,99,102,110]])]:
 settext(orig[listindex],prefix+':');prev=orig[listindex]
 for j,item in enumerate(items):
  name='Report'+prefix.replace(' ','')+str(j+1);bookmark(item,name)
  txt=''.join(item.xpath('.//w:t/text()',namespaces=NS))
  p=para(txt+' — page ');pageref(p,name);prev.addnext(p);prev=p
# Use the original explicit page-break component to give contents its own front-matter page.
orig[21].addprevious(deepcopy(orig[16]))
orig[23].addprevious(deepcopy(orig[16]))

settings=E.fromstring(parts['word/settings.xml']);update=settings.find('w:updateFields',NS)
if update is None:update=E.SubElement(settings,q('updateFields'))
update.set(q('val'),'true');parts['word/settings.xml']=E.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['word/document.xml']=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
 for n,data in parts.items():z.writestr(n,data)
print(OUT)
print('Words:',sum(len(t.split()) for t in tree.xpath('//w:t/text()',namespaces=NS)))
print('Source SHA256:',hashlib.sha256(SOURCE.read_bytes()).hexdigest())
